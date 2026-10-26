"""Build the cheerful eleven-slide presentation for the SMT project."""

# pylint: disable=line-too-long,too-many-arguments,too-many-positional-arguments
# pylint: disable=too-many-locals,unused-argument,import-error,protected-access

from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from smt_changeover.analysis import (
    BLUE,
    CORAL,
    CREAM,
    GREEN,
    INK,
    LIGHT_BLUE,
    MINT,
    TEAL,
    YELLOW,
)


PROJECT_DIR = Path(__file__).resolve().parents[1]
ASSET_DIR = PROJECT_DIR / "assets"
CHART_DIR = PROJECT_DIR / "outputs" / "charts"
PREVIEW_DIR = PROJECT_DIR / "outputs" / "previews"
OUTPUT = PROJECT_DIR / "docs" / "apresentacao" / "apresentacao_smt_em_tempo.pptx"
SUMMARY_PATH = PROJECT_DIR / "outputs" / "metrics" / "resumo_resultados.json"
FONT = "Noto Sans"
TOTAL_SLIDES = 11


def rgb(color: str) -> RGBColor:
    """Convert a hexadecimal color into a PowerPoint color."""
    return RGBColor.from_string(color.lstrip("#"))


def set_background(slide, color: str) -> None:
    """Set a flat slide background."""
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    width: float,
    height: float,
    *,
    size: float = 20,
    color: str = INK,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0.04,
):
    """Add a text box with the project typography."""
    box = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(height)
    )
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.name = FONT
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = rgb(color)
    return box


def add_lines(
    slide,
    lines: list[str],
    x: float,
    y: float,
    width: float,
    height: float,
    *,
    size: float = 17,
    color: str = INK,
    bullet: bool = False,
    spacing: float = 7,
):
    """Add multiple compact lines."""
    box = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(height)
    )
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.06)
    frame.margin_right = Inches(0.06)
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = f"•  {line}" if bullet else line
        paragraph.font.name = FONT
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = rgb(color)
        paragraph.space_after = Pt(spacing)
    return box


def add_bullet_list(
    slide,
    lines: list[str],
    x: float,
    y: float,
    width: float,
    *,
    size: float = 15,
    color: str = INK,
    dot_color: str = CORAL,
    gap: float = 0.48,
) -> None:
    """Add an aligned list with slim contemporary markers."""
    for index, line in enumerate(lines):
        top = y + index * gap
        add_shape(slide, MSO_SHAPE.RECTANGLE, x + 0.03, top + 0.08, 0.06, 0.23, dot_color)
        add_text(
            slide,
            line,
            x + 0.26,
            top,
            width - 0.26,
            gap,
            size=size,
            color=color,
            valign=MSO_ANCHOR.MIDDLE,
            margin=0,
        )


def add_card(
    slide,
    x: float,
    y: float,
    width: float,
    height: float,
    *,
    color: str = "#FFFFFF",
    line: str | None = None,
):
    """Add a clean rounded information card without theme shadows."""
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height),
    )
    card.adjustments[0] = 0.035
    card.fill.solid()
    card.fill.fore_color.rgb = rgb(color)
    if line:
        card.line.color.rgb = rgb(line)
        card.line.width = Pt(1.2)
    else:
        card.line.fill.background()
    card.shadow.inherit = False
    remove_theme_effects(card)
    return card


def add_picture_contain(slide, path: Path, x: float, y: float, width: float, height: float):
    """Place an image inside a box without distortion."""
    with Image.open(path) as image:
        ratio = image.width / image.height
    box_ratio = width / height
    if ratio >= box_ratio:
        draw_width = width
        draw_height = width / ratio
        left = x
        top = y + (height - draw_height) / 2
    else:
        draw_height = height
        draw_width = height * ratio
        left = x + (width - draw_width) / 2
        top = y
    return slide.shapes.add_picture(
        str(path),
        Inches(left),
        Inches(top),
        width=Inches(draw_width),
        height=Inches(draw_height),
    )


def add_picture_opacity(
    slide,
    path: Path,
    x: float,
    y: float,
    width: float,
    height: float,
    opacity: float,
):
    """Place a translucent PNG while preserving transparency."""
    with Image.open(path).convert("RGBA") as image:
        alpha = image.getchannel("A").point(lambda value: int(value * opacity))
        image.putalpha(alpha)
        buffer = BytesIO()
        image.save(buffer, format="PNG")
        buffer.seek(0)
        ratio = image.width / image.height
    box_ratio = width / height
    if ratio >= box_ratio:
        draw_width = width
        draw_height = width / ratio
        left = x
        top = y + (height - draw_height) / 2
    else:
        draw_height = height
        draw_width = height * ratio
        left = x + (width - draw_width) / 2
        top = y
    return slide.shapes.add_picture(
        buffer,
        Inches(left),
        Inches(top),
        width=Inches(draw_width),
        height=Inches(draw_height),
    )


def add_shape(slide, shape_type, x, y, width, height, color, rotation=0):
    """Add one flat decorative shape."""
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.fill.background()
    shape.shadow.inherit = False
    remove_theme_effects(shape)
    shape.rotation = rotation
    return shape


def remove_theme_effects(shape) -> None:
    """Prevent office renderers from restoring theme shadows on a flat shape."""
    for effect_ref in shape._element.xpath("./p:style/a:effectRef"):
        effect_ref.set("idx", "0")


def add_doodles(slide, variant: int, dark: bool = False) -> None:
    """Add one professional illustration as a subtle page background."""
    artwork = (
        ASSET_DIR / "ilustracao-linha-smt.png"
        if variant % 3 == 1
        else ASSET_DIR / "ilustracao-previsao.png"
    )
    opacity = 0.12 if dark else 0.10
    add_picture_opacity(slide, artwork, 5.25, 4.15, 7.75, 3.15, opacity)


def add_header(slide, number: str, title: str, *, color: str = TEAL) -> None:
    """Add a numbered slide header."""
    add_text(slide, number, 0.46, 0.35, 0.48, 0.34, size=12, color=color, bold=True)
    add_text(
        slide,
        title,
        1.10,
        0.31,
        11.12,
        0.60,
        size=24.5,
        color=color,
        bold=True,
        align=PP_ALIGN.CENTER,
    )


def add_footer(slide, current: int, *, color: str = TEAL) -> None:
    """Add the slide counter."""
    add_text(
        slide,
        f"{current:02d} / {TOTAL_SLIDES:02d}",
        11.98,
        7.04,
        0.92,
        0.22,
        size=9.5,
        color=color,
        bold=True,
        align=PP_ALIGN.RIGHT,
        valign=MSO_ANCHOR.MIDDLE,
        margin=0,
    )


def add_metric(slide, value: str, label: str, x: float, y: float, width: float, color: str):
    """Add one compact metric card."""
    add_card(slide, x, y, width, 1.18, color="#FFFFFF", line=MINT)
    add_shape(slide, MSO_SHAPE.RECTANGLE, x + 0.20, y + 0.13, width - 0.40, 0.045, color)
    add_text(
        slide,
        value,
        x + 0.12,
        y + 0.25,
        width - 0.24,
        0.39,
        size=22,
        color=color,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        label,
        x + 0.12,
        y + 0.75,
        width - 0.24,
        0.27,
        size=10.5,
        color=INK,
        align=PP_ALIGN.CENTER,
    )


def slide_cover(prs: Presentation) -> None:
    """Build the presentation cover."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, CREAM)
    add_picture_contain(
        slide,
        ASSET_DIR / "logo-smt-em-tempo.png",
        4.45,
        3.92,
        4.43,
        2.48,
    )
    add_picture_contain(slide, ASSET_DIR / "logo-uea.png", 10.45, 0.23, 1.15, 0.62)
    add_picture_contain(slide, ASSET_DIR / "logo-sialabs.png", 11.72, 0.19, 0.92, 0.68)

    add_text(
        slide,
        "SMT em Tempo",
        1.45,
        1.05,
        10.43,
        0.72,
        size=34,
        color=TEAL,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "Quais trocas de produto precisam de apoio antes de começar?",
        2.05,
        1.86,
        9.23,
        0.85,
        size=24,
        color=INK,
        align=PP_ALIGN.CENTER,
    )
    add_shape(slide, MSO_SHAPE.RECTANGLE, 5.10, 2.78, 3.13, 0.055, TEAL)
    add_text(
        slide,
        "Machine Learning",
        3.47,
        3.10,
        6.40,
        0.36,
        size=18,
        color=TEAL,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "Projeto 4 · Trabalho final · Outubro de 2026",
        3.20,
        3.54,
        6.93,
        0.30,
        size=13.5,
        color=INK,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "Juliana Ballin Lima  ·  Fernanda de Oliveira da Costa  ·  Pedro Henrique Oliveira Dias",
        2.70,
        6.62,
        8.25,
        0.32,
        size=11.5,
        color=TEAL,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, 1)


def slide_summary(prs: Presentation) -> None:
    """Build the agenda slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, TEAL)
    add_doodles(slide, 0, dark=True)
    add_header(slide, "01", "Sumário", color=CREAM)
    labels = [
        ("01", "Problema", "Por que prever", CORAL),
        ("02", "Perguntas", "Número e risco", YELLOW),
        ("03", "Dataset", "2.057 registros", BLUE),
        ("04", "Método", "Limpar e dividir", GREEN),
        ("05", "Resultados", "Regressão e classe", CORAL),
        ("06", "Notebook", "Demonstração", BLUE),
        ("07", "Entregáveis", "Relatório bônus", YELLOW),
        ("08", "Decisão", "Como apoiar", GREEN),
    ]
    for index, (number, title, subtitle, accent) in enumerate(labels):
        column = index % 4
        row = index // 4
        x = 0.70 + column * 3.12
        y = 1.42 + row * 2.18
        add_card(slide, x, y, 2.68, 1.68, color=CREAM, line=MINT)
        add_shape(slide, MSO_SHAPE.RECTANGLE, x + 0.24, y + 0.18, 2.20, 0.045, accent)
        add_text(
            slide,
            number,
            x + 0.30,
            y + 0.36,
            2.08,
            0.20,
            size=9.5,
            color=accent,
            bold=True,
            align=PP_ALIGN.CENTER,
            margin=0,
        )
        add_text(
            slide,
            title,
            x + 0.30,
            y + 0.67,
            2.08,
            0.35,
            size=15,
            color=TEAL,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(
            slide,
            subtitle,
            x + 0.30,
            y + 1.13,
            2.08,
            0.35,
            size=11,
            color=INK,
            align=PP_ALIGN.CENTER,
        )
    add_text(
        slide,
        "Uma previsão para apoiar a decisão da linha",
        2.85,
        6.13,
        7.62,
        0.45,
        size=21,
        color=CREAM,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, 2, color=CREAM)


def slide_problem(prs: Presentation) -> None:
    """Explain the production problem."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, CREAM)
    add_header(slide, "02", "O problema")
    add_card(slide, 0.76, 1.30, 5.18, 4.82, color="#FFFFFF", line=MINT)
    add_text(slide, "A linha para para trocar de produto", 1.12, 1.67, 4.46, 0.50, size=23, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_bullet_list(
        slide,
        [
            "retirar alimentadores do produto anterior",
            "instalar o novo conjunto",
            "carregar o programa da máquina",
            "conferir a primeira placa",
        ],
        1.15,
        2.42,
        4.30,
        size=15.5,
        dot_color=CORAL,
        gap=0.54,
    )
    add_card(slide, 1.12, 4.97, 4.46, 0.69, color=CORAL)
    add_text(slide, "Enquanto isso, a linha não produz.", 1.27, 5.16, 4.16, 0.28, size=15, color="#FFFFFF", bold=True, align=PP_ALIGN.CENTER)

    add_picture_opacity(
        slide,
        ASSET_DIR / "logo-smt-em-tempo.png",
        6.35,
        1.10,
        6.18,
        3.50,
        0.16,
    )
    add_text(slide, "A duração varia muito", 7.10, 1.52, 4.75, 0.40, size=21, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_card(slide, 6.62, 3.24, 2.78, 1.38, color=LIGHT_BLUE, line="#A9D8E8")
    add_text(slide, "20 min", 6.86, 3.53, 2.30, 0.42, size=25, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "trocas rápidas", 6.86, 4.05, 2.30, 0.28, size=12.5, color=INK, align=PP_ALIGN.CENTER)
    add_card(slide, 9.64, 3.24, 2.78, 1.38, color="#F8D2C8", line="#F0B7A9")
    add_text(slide, "> 60 min", 9.88, 3.53, 2.30, 0.42, size=25, color=CORAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "trocas críticas", 9.88, 4.05, 2.30, 0.28, size=12.5, color=INK, align=PP_ALIGN.CENTER)
    add_card(slide, 6.62, 5.06, 5.80, 0.92, color=TEAL)
    add_text(
        slide,
        "A ajuda adicional existe, mas não alcança todas as trocas.",
        6.92,
        5.28,
        5.20,
        0.46,
        size=17,
        color=CREAM,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, 3)


def slide_questions(prs: Presentation) -> None:
    """Show the regression and classification questions."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, CREAM)
    add_doodles(slide, 2)
    add_header(slide, "03", "As perguntas")
    add_text(
        slide,
        "Antes de a linha parar, quais trocas precisam de mais atenção?",
        1.48,
        1.20,
        10.35,
        0.62,
        size=25,
        color=INK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_card(slide, 0.78, 2.30, 5.70, 3.28, color=LIGHT_BLUE, line="#A9D8E8")
    add_shape(slide, MSO_SHAPE.RECTANGLE, 1.18, 2.66, 4.90, 0.055, BLUE)
    add_text(slide, "01", 1.18, 2.90, 4.90, 0.25, size=10.5, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Regressão", 1.18, 3.20, 4.90, 0.42, size=24, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Quantos minutos a troca deve levar?", 1.23, 3.87, 4.85, 0.70, size=20, color=INK, align=PP_ALIGN.CENTER)
    add_text(slide, "Métrica principal: MAE", 1.55, 4.86, 4.20, 0.34, size=14, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_card(slide, 6.86, 2.30, 5.70, 3.28, color="#F8D2C8", line="#F0B7A9")
    add_shape(slide, MSO_SHAPE.RECTANGLE, 7.26, 2.66, 4.90, 0.055, CORAL)
    add_text(slide, "02", 7.26, 2.90, 4.90, 0.25, size=10.5, color=CORAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Classificação", 7.26, 3.20, 4.90, 0.42, size=24, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "A troca deve passar de 42 minutos?", 7.31, 3.87, 4.85, 0.70, size=20, color=INK, align=PP_ALIGN.CENTER)
    add_text(slide, "Foco: encontrar o risco sem alarmar demais", 7.36, 4.79, 4.75, 0.42, size=13.5, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 4)


def slide_dataset(prs: Presentation, summary: dict) -> None:
    """Describe data scope, features, and temporal split."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, TEAL)
    add_doodles(slide, 0, dark=True)
    add_header(slide, "04", "O dataset", color=CREAM)
    add_metric(slide, "2.057", "linhas recebidas", 0.66, 1.18, 2.18, CORAL)
    add_metric(slide, "2.043", "trocas após deduplicar", 3.02, 1.18, 2.36, GREEN)
    add_metric(slide, "27,2%", "acima de 42 min", 5.56, 1.18, 2.18, CORAL)
    add_metric(slide, "8", "entradas antes da troca", 7.92, 1.18, 2.30, BLUE)
    add_metric(slide, "jan a ago", "período de 2026", 10.40, 1.18, 2.25, YELLOW)

    add_card(slide, 0.66, 2.73, 5.85, 3.38, color=CREAM, line=MINT)
    add_text(slide, "Entradas conhecidas antes", 1.00, 3.02, 5.05, 0.42, size=18, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_bullet_list(
        slide,
        [
            "alimentadores e kit preparado",
            "experiência e componentes novos",
            "tamanho do lote",
            "linha, mudança de família e turno",
        ],
        1.02,
        3.66,
        5.05,
        size=15,
        dot_color=GREEN,
        gap=0.49,
    )

    add_card(slide, 6.82, 2.73, 5.85, 3.38, color=CREAM, line=MINT)
    add_text(slide, "Divisão temporal", 7.17, 3.02, 4.95, 0.42, size=18, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    split_details = [
        ("Treino", "1.236", "05/01 a 27/05", GREEN),
        ("Validação", "404", "28/05 a 15/07", YELLOW),
        ("Teste", "403", "16/07 a 28/08", CORAL),
    ]
    for index, (name, amount, period, color) in enumerate(split_details):
        y = 3.72 + index * 0.68
        add_shape(slide, MSO_SHAPE.RECTANGLE, 7.22, y, 0.07, 0.26, color)
        add_text(slide, name, 7.53, y - 0.04, 1.40, 0.30, size=13.5, color=INK, bold=True)
        add_text(slide, amount, 9.03, y - 0.04, 0.72, 0.30, size=13.5, color=TEAL, bold=True)
        add_text(slide, period, 9.88, y - 0.04, 2.05, 0.30, size=12.5, color=INK)
    add_text(slide, "Teste reservado até o final", 7.18, 5.61, 4.95, 0.30, size=13, color=CORAL, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 5, color=CREAM)


def slide_method(prs: Presentation) -> None:
    """Summarize data cleaning and exploratory evidence."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, CREAM)
    add_doodles(slide, 1)
    add_header(slide, "05", "Limpeza e relações principais")
    add_picture_contain(slide, CHART_DIR / "02_relacoes_principais.png", 0.55, 1.10, 7.15, 4.90)
    add_card(slide, 7.90, 1.18, 4.82, 2.11, color=MINT)
    add_text(slide, "O que foi tratado", 8.22, 1.47, 4.15, 0.36, size=18, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_bullet_list(
        slide,
        ["14 duplicatas", "grafias e datas", "vazios e código 9999", "medidas impossíveis"],
        8.20,
        1.96,
        4.10,
        size=12.8,
        dot_color=GREEN,
        gap=0.30,
    )
    add_card(slide, 7.90, 3.54, 4.82, 2.43, color="#F8D2C8")
    add_text(slide, "O que os gráficos sugerem", 8.22, 3.84, 4.15, 0.36, size=18, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_bullet_list(
        slide,
        ["kit preparado encurta a troca", "mais alimentadores alongam a troca", "associação não prova causa"],
        8.20,
        4.38,
        4.10,
        size=13.2,
        dot_color=CORAL,
        gap=0.39,
    )
    add_card(slide, 1.45, 6.18, 10.42, 0.55, color=TEAL)
    add_text(slide, "A coluna posterior ao evento ficou fora dos modelos.", 1.72, 6.31, 9.88, 0.28, size=15, color=CREAM, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 6)


def slide_regression(prs: Presentation, summary: dict) -> None:
    """Present the final regression result."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, CREAM)
    add_doodles(slide, 2)
    add_header(slide, "06", "Resultado 1 · Quantos minutos?")
    add_picture_contain(slide, CHART_DIR / "04_comparacao_regressao.png", 0.46, 1.23, 7.23, 4.95)
    metric_rows = [
        (1.33, "2,66 min", "erro médio da árvore", TEAL),
        (2.83, "67,4%", "menos erro que a referência", CORAL),
        (4.33, "R² 0,884", "no período de teste", BLUE),
    ]
    for y, value, label, accent in metric_rows:
        add_card(slide, 7.88, y, 4.87, 1.24, color="#FFFFFF", line=MINT)
        add_shape(slide, MSO_SHAPE.RECTANGLE, 8.16, y + 0.16, 4.31, 0.045, accent)
        add_text(slide, value, 8.16, y + 0.30, 4.31, 0.42, size=25, color=accent, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, label, 8.16, y + 0.82, 4.31, 0.25, size=12.5, color=INK, align=PP_ALIGN.CENTER)
    add_card(slide, 1.43, 6.17, 10.43, 0.55, color=TEAL)
    add_text(slide, "A árvore de profundidade 7 foi escolhida na validação.", 1.68, 6.30, 9.92, 0.28, size=15, color=CREAM, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 7)


def slide_classification(prs: Presentation, summary: dict) -> None:
    """Present the final classification result."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, CREAM)
    add_doodles(slide, 0)
    add_header(slide, "07", "Resultado 2 · Vai passar de 42 minutos?")
    add_picture_contain(slide, CHART_DIR / "06_matriz_confusao.png", 0.65, 1.18, 6.35, 5.28)
    classification_rows = [
        (1.29, "89", "trocas demoradas encontradas", TEAL, MINT),
        (2.60, "13", "trocas demoradas perdidas", CORAL, "#F8D2C8"),
        (3.91, "31", "falsos alarmes", BLUE, LIGHT_BLUE),
    ]
    for y, value, label, accent, fill in classification_rows:
        add_card(slide, 7.16, y, 5.49, 1.09, color=fill, line=MINT)
        add_shape(slide, MSO_SHAPE.RECTANGLE, 7.46, y + 0.14, 4.89, 0.045, accent)
        add_text(slide, value, 7.46, y + 0.25, 4.89, 0.34, size=24, color=accent, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, label, 7.46, y + 0.68, 4.89, 0.25, size=12.8, color=INK, align=PP_ALIGN.CENTER)
    add_card(slide, 7.16, 5.22, 5.49, 1.05, color=TEAL)
    add_text(slide, "Recall 87,3%  ·  Precisão 74,2%  ·  F1 80,2%", 7.43, 5.55, 4.95, 0.30, size=13.2, color=CREAM, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 8)


def slide_notebook(prs: Presentation) -> None:
    """Show the executed notebook."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, TEAL)
    add_doodles(slide, 1, dark=True)
    add_header(slide, "08", "Demonstração do notebook", color=CREAM)
    preview = PREVIEW_DIR / "preview-notebook.png"
    add_card(slide, 0.62, 1.10, 8.05, 5.67, color="#FFFFFF")
    add_picture_contain(slide, preview, 0.74, 1.23, 7.81, 5.42)
    callouts = [
        ("1", "Limpeza", "cada problema é observado e conferido", CORAL),
        ("2", "Modelos", "referência, regressão linear e árvores", YELLOW),
        ("3", "Entrega", "células executadas e resultados reproduzíveis", BLUE),
    ]
    for index, (number, title, body, color) in enumerate(callouts):
        y = 1.30 + index * 1.67
        add_card(slide, 8.94, y, 3.77, 1.35, color=CREAM, line=MINT)
        add_shape(slide, MSO_SHAPE.RECTANGLE, 9.22, y + 0.17, 3.21, 0.045, color)
        add_text(slide, f"0{number}", 9.22, y + 0.30, 3.21, 0.22, size=9.5, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, title, 9.22, y + 0.55, 3.21, 0.30, size=15, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, body, 9.22, y + 0.88, 3.21, 0.34, size=11.2, color=INK, align=PP_ALIGN.CENTER)
    add_footer(slide, 9, color=CREAM)


def slide_deliverables(prs: Presentation) -> None:
    """Present the notebook, slides, and bonus technical report."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, CREAM)
    add_doodles(slide, 1)
    add_header(slide, "09", "Entregáveis")

    card_specs = [
        (0.70, "Notebook", "entrega principal", BLUE),
        (4.57, "Apresentação", "PPTX e PDF", GREEN),
        (8.44, "Relatório técnico", "material complementar", CORAL),
    ]
    for x, title, subtitle, accent in card_specs:
        add_card(slide, x, 1.26, 3.52, 5.20, color="#FFFFFF", line=MINT)
        add_shape(slide, MSO_SHAPE.RECTANGLE, x + 0.34, 1.55, 2.84, 0.055, accent)
        add_text(
            slide,
            title,
            x + 0.28,
            1.86,
            2.96,
            0.40,
            size=18,
            color=TEAL,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(
            slide,
            subtitle,
            x + 0.34,
            2.37,
            2.84,
            0.34,
            size=11.5,
            color=INK,
            align=PP_ALIGN.CENTER,
        )

    add_picture_contain(slide, PREVIEW_DIR / "preview-notebook.png", 1.02, 3.02, 2.88, 2.36)
    add_picture_opacity(slide, ASSET_DIR / "logo-smt-em-tempo.png", 5.12, 2.92, 2.42, 2.42, 0.48)
    add_picture_contain(slide, PREVIEW_DIR / "preview-relatorio.png", 9.43, 2.89, 1.54, 2.18)

    add_card(slide, 9.36, 5.48, 1.70, 0.44, color=YELLOW)
    add_text(
        slide,
        "BÔNUS",
        9.36,
        5.58,
        1.70,
        0.20,
        size=10.5,
        color=INK,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0,
    )
    add_text(
        slide,
        "Todos os arquivos estão organizados no repositório.",
        2.70,
        6.72,
        7.90,
        0.28,
        size=13,
        color=TEAL,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, 10)


def slide_conclusion(prs: Presentation) -> None:
    """Close with the operational decision and limitations."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, CREAM)
    add_doodles(slide, 2)
    add_header(slide, "10", "Decisão apoiada e aprendizados")
    add_card(slide, 0.72, 1.18, 7.14, 4.92, color=TEAL)
    add_text(slide, "Como usar", 1.12, 1.55, 6.34, 0.38, size=21, color=CREAM, bold=True, align=PP_ALIGN.CENTER)
    steps = [
        ("1", "Calcular o risco antes da troca"),
        ("2", "Ordenar as trocas mais arriscadas"),
        ("3", "Direcionar apoio com decisão humana"),
        ("4", "Monitorar perdas e falsos alarmes"),
    ]
    for index, (number, label) in enumerate(steps):
        y = 2.20 + index * 0.80
        accent = [CORAL, YELLOW, BLUE, GREEN][index]
        add_shape(slide, MSO_SHAPE.RECTANGLE, 1.20, y + 0.05, 0.07, 0.34, accent)
        add_text(slide, f"0{number}", 1.46, y + 0.08, 0.48, 0.22, size=10.5, color=accent, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, label, 2.08, y + 0.04, 5.02, 0.32, size=15, color=CREAM, bold=index == 2)
    add_card(slide, 8.15, 1.18, 4.46, 2.15, color=MINT)
    add_text(slide, "Pontos de aprendizado", 8.48, 1.48, 3.75, 0.38, size=18, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_bullet_list(
        slide,
        ["avaliar no tempo", "comparar com referência", "evitar vazamento"],
        8.48,
        2.03,
        3.65,
        size=13.4,
        dot_color=GREEN,
        gap=0.34,
    )
    add_card(slide, 8.15, 3.58, 4.46, 2.52, color="#F8D2C8")
    add_text(slide, "Limites", 8.48, 3.88, 3.75, 0.38, size=18, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_bullet_list(
        slide,
        ["dados sintéticos", "sem prova causal", "validar antes do uso real"],
        8.48,
        4.43,
        3.65,
        size=13.4,
        dot_color=CORAL,
        gap=0.38,
    )
    add_text(slide, "Prever cedo para apoiar melhor.", 3.16, 6.32, 7.05, 0.40, size=21, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 11)


def build_presentation() -> None:
    """Create and save the full deck."""
    summary = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    slide_cover(presentation)
    slide_summary(presentation)
    slide_problem(presentation)
    slide_questions(presentation)
    slide_dataset(presentation, summary)
    slide_method(presentation)
    slide_regression(presentation, summary)
    slide_classification(presentation, summary)
    slide_notebook(presentation)
    slide_deliverables(presentation)
    slide_conclusion(presentation)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(OUTPUT)
    print(f"Apresentação criada em: {OUTPUT}")


if __name__ == "__main__":
    build_presentation()
